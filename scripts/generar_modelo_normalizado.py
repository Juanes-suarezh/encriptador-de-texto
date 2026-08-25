#!/usr/bin/env python3
"""Modelo normalizado (15 entidades): diagramas, drawio y Word de la actividad 2."""
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon, Rectangle
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path("/workspace/docs")
FIG = ROOT / "figuras"
DIA = ROOT / "diagramas"
FIG.mkdir(exist_ok=True)
DIA.mkdir(exist_ok=True)

GREEN, GREEN_FILL = "#1B4D3E", "#E8F2EB"
GOLD, GOLD_FILL = "#8A6A2F", "#F8EBD0"
NAVY, NAVY_FILL = "#2C4A6E", "#E7EEF6"
INK, LINE = "#1A1A1A", "#4A4A4A"


# ---------------------------------------------------------------------------
# Chen PNG
# ---------------------------------------------------------------------------
def box(ax, x, y, text, w=2.4, h=0.62, fc=GREEN_FILL, ec=GREEN, fs=9):
    ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h,
                                boxstyle="round,pad=0.02,rounding_size=0.08",
                                lw=1.5, ec=ec, fc=fc, zorder=3))
    ax.text(x, y, text, ha="center", va="center", fontsize=fs, fontweight="bold",
            color=INK, zorder=4, fontfamily="serif")
    return x, y, w, h


def dia(ax, x, y, text, w=1.45, h=0.52, fs=7):
    pts = np.array([[x, y + h / 2], [x + w / 2, y], [x, y - h / 2], [x - w / 2, y]])
    ax.add_patch(Polygon(pts, closed=True, fc=GOLD_FILL, ec=GOLD, lw=1.2, zorder=3))
    ax.text(x, y, text, ha="center", va="center", fontsize=fs, color=INK, zorder=4, fontfamily="serif")
    return x, y


def hline(ax, x1, x2, y):
    ax.plot([x1, x2], [y, y], color=LINE, lw=1.1, zorder=1)


def vline(ax, x, y1, y2):
    ax.plot([x, x], [y1, y2], color=LINE, lw=1.1, zorder=1)


def card(ax, x, y, t):
    ax.text(x, y, t, fontsize=8, color=GREEN, fontweight="bold", ha="center", va="center",
            zorder=5, fontfamily="serif",
            bbox=dict(boxstyle="round,pad=0.08", fc="white", ec="none", alpha=0.95))


def draw_chen():
    fig, ax = plt.subplots(figsize=(15.2, 9.6), dpi=220)
    ax.set_xlim(0, 22)
    ax.set_ylim(0, 13.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_title("DER normalizado (15 entidades) — notación de Chen\n"
                 "Sin atributos derivados ni tablas que repetían las mismas llaves",
                 fontsize=13, fontweight="bold", fontfamily="serif", color=GREEN, pad=8)

    ax.add_patch(Rectangle((0.25, 11.15), 21.5, 1.55, fc="#FAFBF9", ec="#D7DED8", lw=0.6, zorder=0))
    ax.add_patch(Rectangle((0.25, 8.35), 21.5, 2.65, fc="#FAFBF9", ec="#D7DED8", lw=0.6, zorder=0))
    ax.add_patch(Rectangle((0.25, 5.15), 21.5, 3.05, fc="#FAFBF9", ec="#D7DED8", lw=0.6, zorder=0))
    ax.add_patch(Rectangle((0.25, 2.55), 21.5, 2.45, fc="#FAFBF9", ec="#D7DED8", lw=0.6, zorder=0))
    ax.add_patch(Rectangle((0.25, 0.35), 21.5, 2.05, fc="#FAFBF9", ec="#D7DED8", lw=0.6, zorder=0))
    ax.text(0.4, 12.45, "Actores", fontsize=8.5, color=GOLD, fontweight="bold", fontfamily="serif")
    ax.text(0.4, 10.75, "Maestros (una sola vez cada hecho de negocio)", fontsize=8.5, color=GOLD, fontweight="bold", fontfamily="serif")
    ax.text(0.4, 7.95, "Estructura tarifaria", fontsize=8.5, color=GOLD, fontweight="bold", fontfamily="serif")
    ax.text(0.4, 4.75, "Proceso (un archivo, una validación)", fontsize=8.5, color=GOLD, fontweight="bold", fontfamily="serif")
    ax.text(0.4, 2.15, "Reforma (una sola entidad con tipo)", fontsize=8.5, color=GOLD, fontweight="bold", fontfamily="serif")

    u = box(ax, 11, 11.75, "USUARIO", w=2.6)
    c = box(ax, 2.6, 9.55, "CLIENTE")
    d_t = dia(ax, 4.85, 9.55, "tiene")
    ce = box(ax, 7.1, 9.55, "CENTRO")
    hline(ax, c[0] + c[2] / 2, d_t[0] - 0.72, 9.55)
    hline(ax, d_t[0] + 0.72, ce[0] - ce[2] / 2, 9.55)
    card(ax, 3.7, 9.85, "1")
    card(ax, 6.0, 9.85, "N")
    p = box(ax, 11.0, 9.55, "PRODUCTO", w=2.6)
    pa = box(ax, 15.2, 9.55, "PAÍS", w=2.2)
    d_c = dia(ax, 17.35, 9.55, "contiene")
    cp = box(ax, 19.7, 9.55, "CÓDIGO_POSTAL", w=2.8, fs=8.2)
    hline(ax, pa[0] + pa[2] / 2, d_c[0] - 0.72, 9.55)
    hline(ax, d_c[0] + 0.72, cp[0] - cp[2] / 2, 9.55)
    card(ax, 16.25, 9.85, "1")
    card(ax, 18.4, 9.85, "N")

    t = box(ax, 2.6, 6.55, "TARIFA")
    d_tv = dia(ax, 4.85, 6.55, "tiene")
    v = box(ax, 7.1, 6.55, "VIGENCIA")
    d_def = dia(ax, 9.4, 6.55, "define")
    b = box(ax, 11.7, 6.55, "BAREMO")
    d_des = dia(ax, 14.15, 6.55, "desglosa")
    db = box(ax, 16.7, 6.55, "DETALLE_BAREMO", w=2.9, fs=8.2)
    ru = box(ax, 20.0, 6.55, "RUTA", w=2.2)
    for a, d, e, la, lb, xa, xb in [
        (t, d_tv, v, "1", "N", 3.7, 6.0),
        (v, d_def, b, "1", "N", 8.25, 10.5),
        (b, d_des, db, "1", "N", 12.9, 15.5),
    ]:
        hline(ax, a[0] + a[2] / 2, d[0] - 0.72, 6.55)
        hline(ax, d[0] + 0.72, e[0] - e[2] / 2, 6.55)
        card(ax, xa, 6.85, la)
        card(ax, xb, 6.85, lb)

    d_pos = dia(ax, 2.6, 8.05, "posee")
    vline(ax, 2.6, c[1] - 0.31, d_pos[1] + 0.26)
    vline(ax, 2.6, d_pos[1] - 0.26, t[1] + 0.31)
    card(ax, 3.05, 8.75, "1")
    card(ax, 3.05, 7.35, "N")

    d_usa = dia(ax, 11.0, 8.05, "usa")
    vline(ax, 11.0, p[1] - 0.31, d_usa[1] + 0.26)
    vline(ax, 11.0, d_usa[1] - 0.26, b[1] + 0.31)
    card(ax, 11.45, 8.75, "1")
    card(ax, 11.45, 7.35, "N")

    d_cub = dia(ax, 7.1, 5.55, "cubre")
    vline(ax, 7.1, v[1] - 0.31, d_cub[1] + 0.26)
    hline(ax, 7.1, 20.0, 5.28)
    vline(ax, 20.0, 5.28, ru[1] - 0.31)
    card(ax, 6.65, 5.95, "1")
    card(ax, 19.5, 5.55, "N")

    d_apl = dia(ax, 16.7, 5.55, "aplica")
    hline(ax, b[0] + 0.4, d_apl[0] - 0.72, 5.55)
    hline(ax, d_apl[0] + 0.72, ru[0] - ru[2] / 2, 5.55)
    card(ax, 14.4, 5.82, "1")
    card(ax, 18.4, 5.82, "N")

    d_ubi = dia(ax, 19.7, 8.05, "ubica")
    vline(ax, 19.7, cp[1] - 0.31, d_ubi[1] + 0.26)
    vline(ax, 19.7, d_ubi[1] - 0.26, ru[1] + 0.31)
    card(ax, 20.2, 8.75, "1")
    card(ax, 20.2, 7.35, "N")

    s = box(ax, 4.2, 3.65, "SOLICITUD", fc=NAVY_FILL, ec=NAVY)
    d_adj = dia(ax, 7.0, 3.65, "genera")
    a = box(ax, 10.0, 3.65, "ARCHIVO", w=2.6, fc=NAVY_FILL, ec=NAVY)
    d_val = dia(ax, 13.1, 3.65, "se valida")
    va = box(ax, 16.4, 3.65, "VALIDACIÓN", w=2.7, fc=NAVY_FILL, ec=NAVY)
    hline(ax, s[0] + s[2] / 2, d_adj[0] - 0.72, 3.65)
    hline(ax, d_adj[0] + 0.72, a[0] - a[2] / 2, 3.65)
    hline(ax, a[0] + a[2] / 2, d_val[0] - 0.72, 3.65)
    hline(ax, d_val[0] + 0.72, va[0] - va[2] / 2, 3.65)
    card(ax, 5.55, 3.95, "1")
    card(ax, 8.35, 3.95, "N")
    card(ax, 11.55, 3.95, "1")
    card(ax, 14.7, 3.95, "N")

    d_reg = dia(ax, 11.0, 10.55, "registra")
    # usuario registra solicitud
    ax.annotate("", xy=(s[0], s[1] + 0.31), xytext=(u[0] - 1.2, u[1] - 0.31),
                arrowprops=dict(arrowstyle="-", color=LINE, lw=1.05), zorder=1)
    card(ax, 8.2, 10.2, "1 : N  registra")

    d_ori = dia(ax, 2.6, 4.55, "origina")
    vline(ax, 2.6, t[1] - 0.31, 4.8)
    vline(ax, 2.6, 4.3, s[1] + 0.31)
    hline(ax, 2.6, s[0] - s[2] / 2, s[1])
    card(ax, 2.15, 5.5, "1")
    card(ax, 3.2, 4.15, "N")

    cr = box(ax, 11.0, 1.2, "CONDICIÓN_REFORMA", w=3.6, fs=8.5, fc=NAVY_FILL, ec=NAVY)
    d_ap = dia(ax, 11.0, 2.05, "aplica")
    vline(ax, 11.0, d_ap[1] - 0.26, cr[1] + 0.31)
    ax.text(11.0, 2.42, "CLIENTE / CENTRO / PRODUCTO   1  →  N", fontsize=7.5,
            ha="center", color=GREEN, fontfamily="serif")

    box(ax, 1.7, 0.55, "Entidad", w=1.7, h=0.32, fs=7)
    dia(ax, 3.7, 0.55, "Relación", w=1.4, h=0.36, fs=6.5)
    ax.text(8.5, 0.55, "1 = uno    N = muchos     País y moneda se leen de maestros, no se copian en RUTA",
            fontsize=7.6, va="center", color=INK, fontfamily="serif")

    fig.tight_layout()
    path = FIG / "diagrama_er_chen.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()
    print("chen", path)


def entity_table(ax, x, y, title, rows, header="#1B4D3E", w=2.55):
    row_h, n = 0.23, len(rows) + 1
    h = n * row_h
    ax.add_patch(plt.Rectangle((x, y - h), w, h, fc="white", ec=header, lw=0.9, zorder=2))
    ax.add_patch(plt.Rectangle((x, y - row_h), w, row_h, fc=header, ec=header, zorder=3))
    ax.text(x + w / 2, y - row_h / 2, title, ha="center", va="center", fontsize=7.4,
            color="white", fontweight="bold", fontfamily="serif", zorder=4)
    for i, row in enumerate(rows):
        yy = y - row_h * (i + 1.5)
        if i % 2 == 0:
            ax.add_patch(plt.Rectangle((x, y - row_h * (i + 2)), w, row_h, fc="#F4F7F4", ec="none", zorder=3))
        ax.text(x + 0.08, yy, row, ha="left", va="center", fontsize=6.3, color=INK, fontfamily="serif", zorder=4)
    return x, y, w, h


def draw_crow():
    fig, ax = plt.subplots(figsize=(16.2, 10.2), dpi=210)
    ax.set_xlim(0, 23)
    ax.set_ylim(0, 14.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_title("Modelo relacional normalizado (MR) — PK / FK, sin atributos copiados",
                 fontsize=13, fontweight="bold", fontfamily="serif", color=GREEN, pad=8)

    g, o, n = "#1B4D3E", "#8A6A2F", "#3D5A80"
    t = {}
    t["USUARIO"] = entity_table(ax, 0.3, 13.7, "USUARIO", ["PK id_usuario", "nombres", "apellidos", "correo UK", "rol", "estado"], g)
    t["CLIENTE"] = entity_table(ax, 0.3, 11.5, "CLIENTE", ["PK id_cliente", "agrupador UK", "nit_oracle", "nombre_cliente", "tipo_cliente", "ciudad", "direccion", "ejecutivo", "fecha_reforma", "estado"], g)
    t["CENTRO"] = entity_table(ax, 0.3, 7.7, "CENTRO", ["PK id_centro", "FK id_cliente", "codigo_centro", "nombre_centro"], g)
    t["PRODUCTO"] = entity_table(ax, 3.3, 13.7, "PRODUCTO", ["PK id_producto", "codigo_producto UK", "nombre_producto", "tipo_producto", "tipo_cliente", "tipo_baremo", "factor_volumetrico", "moneda"], g)
    t["PAIS"] = entity_table(ax, 3.3, 10.3, "PAIS", ["PK codigo_pais", "nombre_pais"], g)
    t["CP"] = entity_table(ax, 3.3, 8.7, "CODIGO_POSTAL", ["PK id_codigo_postal", "FK codigo_pais", "codigo_postal UK", "ciudad"], g)
    t["TARIFA"] = entity_table(ax, 6.3, 13.7, "TARIFA", ["PK id_tarifa", "FK id_cliente", "codigo_tarifa", "tipo_tarifa", "estado"], g)
    t["VIGENCIA"] = entity_table(ax, 6.3, 11.5, "VIGENCIA", ["PK id_vigencia", "FK id_tarifa", "fecha_desde", "fecha_hasta", "estado"], g)
    t["BAREMO"] = entity_table(ax, 9.3, 13.7, "BAREMO", ["PK id_baremo", "FK id_vigencia", "FK id_producto", "numero_baremo", "valor_fijo", "valor_minimo"], g)
    t["DET"] = entity_table(ax, 9.3, 11.3, "DETALLE_BAREMO", ["PK id_detalle", "FK id_baremo", "hasta_kg", "importe", "sobrepeso"], g)
    t["RUTA"] = entity_table(ax, 9.3, 9.15, "RUTA", ["PK id_ruta", "FK id_vigencia", "FK id_producto", "FK id_baremo", "FK id_cp_origen", "FK id_cp_destino"], g)
    t["SOL"] = entity_table(ax, 12.5, 13.7, "SOLICITUD", ["PK id_solicitud", "FK id_cliente", "FK id_usuario", "tipo_solicitud", "canal", "fecha_solicitud", "estado", "observaciones"], o, 2.85)
    t["ARC"] = entity_table(ax, 12.5, 10.55, "ARCHIVO", ["PK id_archivo", "FK id_solicitud", "tipo INSUMO|CSV", "nombre_archivo", "ruta", "hash_integridad", "tipo_fichero", "separador", "FK id_usuario_carga", "FK id_usuario_aprueba", "fecha_cargue", "medio_cargue", "resultado_cargue", "registro_sox"], o, 2.85)
    t["VAL"] = entity_table(ax, 16.0, 13.7, "VALIDACION", ["PK id_validacion", "FK id_archivo", "nombre_campo", "fila", "resultado", "detalle"], o, 2.85)
    t["REF"] = entity_table(ax, 16.0, 11.15, "CONDICION_REFORMA", ["PK id_condicion", "tipo DESCUENTO|MINIMA|CARGO", "FK id_cliente", "FK id_centro", "FK id_producto", "codigo_crr", "codigo_zona", "FK id_cp_origen", "FK id_cp_destino", "fecha_desde / hasta / baja", "concepto_facturable", "desde_kg / hasta_kg", "descuento", "rango_kg", "minimo", "tasa_manejo", "numero_baremo"], n, 3.3)

    def link(a, b):
        x1, y1 = a[0] + a[2], a[1] - a[3] / 2
        x2, y2 = b[0], b[1] - b[3] / 2
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="-|>", color=GREEN, lw=1.05, mutation_scale=8), zorder=1)

    for a, b in [("USUARIO", "SOL"), ("CLIENTE", "CENTRO"), ("CLIENTE", "TARIFA"), ("CLIENTE", "SOL"),
                 ("TARIFA", "VIGENCIA"), ("VIGENCIA", "BAREMO"), ("VIGENCIA", "RUTA"),
                 ("PRODUCTO", "BAREMO"), ("PRODUCTO", "RUTA"), ("BAREMO", "DET"), ("BAREMO", "RUTA"),
                 ("PAIS", "CP"), ("CP", "RUTA"), ("SOL", "ARC"), ("ARC", "VAL"),
                 ("CLIENTE", "REF"), ("PRODUCTO", "REF"), ("CENTRO", "REF")]:
        link(t[a], t[b])

    ax.text(11.5, 0.35, "PK = primaria    FK = foránea    UK = única    → lado N.  "
            "Rol, zona y CRR son atributos (no tablas).  País, moneda y volumétrico se consultan en PRODUCTO/PAÍS.",
            ha="center", fontsize=7.4, color=INK, fontfamily="serif")
    fig.tight_layout()
    path = FIG / "diagrama_er_crowsfoot.png"
    fig.savefig(path, dpi=210, bbox_inches="tight", facecolor="white")
    plt.close()
    print("crow", path)


# ---------------------------------------------------------------------------
# Draw.io
# ---------------------------------------------------------------------------
E = ("rounded=0;whiteSpace=wrap;html=1;fillColor=#E8F0EA;strokeColor=#1B4D3E;"
     "fontStyle=1;fontSize=12;fontFamily=Times New Roman;")
EP = ("rounded=0;whiteSpace=wrap;html=1;fillColor=#E7EEF6;strokeColor=#2C4A6E;"
      "fontStyle=1;fontSize=12;fontFamily=Times New Roman;")
D = ("rhombus;whiteSpace=wrap;html=1;fillColor=#F7E7C6;strokeColor=#8A6A2F;"
     "fontStyle=1;fontSize=11;fontFamily=Times New Roman;")
ED = "endArrow=none;html=1;strokeColor=#555555;strokeWidth=1.2;"
LB = ("text;html=1;strokeColor=none;fillColor=none;align=center;fontStyle=1;"
      "fontSize=11;fontColor=#1B4D3E;fontFamily=Times New Roman;")
TT = ("text;html=1;strokeColor=none;fillColor=none;align=center;fontStyle=1;"
      "fontSize=18;fontColor=#1B4D3E;fontFamily=Times New Roman;")
SW = ("swimlane;fontStyle=1;childLayout=stackLayout;horizontal=1;startSize=28;"
      "horizontalStack=0;resizeParent=1;resizeLast=0;collapsible=0;marginBottom=0;"
      "whiteSpace=wrap;html=1;fillColor={f};strokeColor={f};fontColor=#FFFFFF;"
      "fontSize=12;fontFamily=Times New Roman;")
RW = ("text;strokeColor=none;fillColor={bg};align=left;verticalAlign=middle;spacingLeft=8;"
      "overflow=hidden;rotatable=0;points=[[0,0.5],[1,0.5]];portConstraint=eastwest;"
      "whiteSpace=wrap;html=1;fontSize=10;fontFamily=Times New Roman;{ex}")
RE = "html=1;strokeColor=#1B4D3E;strokeWidth=1.2;startArrow=ERone;startFill=0;endArrow=ERmany;endFill=0;"


def mx_cell(cid, value, style, x, y, w, h, parent="1"):
    return (f'        <mxCell id="{escape(cid)}" value="{escape(value)}" style="{style}" vertex="1" parent="{parent}">\n'
            f'          <mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry" />\n'
            f"        </mxCell>\n")


def mx_edge(cid, src, tgt, style=ED):
    return (f'        <mxCell id="{escape(cid)}" style="{style}" edge="1" parent="1" source="{escape(src)}" target="{escape(tgt)}">\n'
            f'          <mxGeometry relative="1" as="geometry" />\n'
            f"        </mxCell>\n")


def mxfile(name, did, w, h, body):
    return f"""<mxfile host="app.diagrams.net" modified="2026-08-25T23:00:00.000Z" agent="Cursor" version="22.1.0" type="device">
  <diagram id="{did}" name="{escape(name)}">
    <mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="{w}" pageHeight="{h}" math="0" shadow="0">
      <root>
        <mxCell id="0" />
        <mxCell id="1" parent="0" />
{body}      </root>
    </mxGraphModel>
  </diagram>
</mxfile>
"""


def build_der_drawio():
    b = [mx_cell("title", "DER normalizado — 15 entidades (Chen)", TT, 200, 16, 1200, 36)]
    b.append(mx_cell("USUARIO", "USUARIO", E, 620, 70, 150, 48))
    b.append(mx_cell("CLIENTE", "CLIENTE", E, 40, 200, 150, 48))
    b.append(mx_cell("d_tiene", "tiene", D, 220, 196, 100, 54))
    b.append(mx_cell("CENTRO", "CENTRO", E, 350, 200, 140, 48))
    b.append(mx_cell("PRODUCTO", "PRODUCTO", E, 560, 200, 160, 48))
    b.append(mx_cell("PAIS", "PAÍS", E, 800, 200, 120, 48))
    b.append(mx_cell("d_cont", "contiene", D, 950, 196, 110, 54))
    b.append(mx_cell("CP", "CÓDIGO_POSTAL", E, 1090, 200, 170, 48))
    b += [mx_edge("e1", "CLIENTE", "d_tiene"), mx_edge("e2", "d_tiene", "CENTRO"),
          mx_edge("e3", "PAIS", "d_cont"), mx_edge("e4", "d_cont", "CP")]
    b.append(mx_cell("c1", "1", LB, 200, 180, 20, 18))
    b.append(mx_cell("c1n", "N", LB, 330, 180, 20, 18))
    b.append(mx_cell("c2", "1", LB, 930, 180, 20, 18))
    b.append(mx_cell("c2n", "N", LB, 1070, 180, 20, 18))

    b.append(mx_cell("TARIFA", "TARIFA", E, 40, 400, 150, 48))
    b.append(mx_cell("d_tv", "tiene", D, 220, 396, 100, 54))
    b.append(mx_cell("VIGENCIA", "VIGENCIA", E, 350, 400, 150, 48))
    b.append(mx_cell("d_def", "define", D, 530, 396, 100, 54))
    b.append(mx_cell("BAREMO", "BAREMO", E, 660, 400, 140, 48))
    b.append(mx_cell("d_des", "desglosa", D, 830, 396, 110, 54))
    b.append(mx_cell("DETALLE", "DETALLE_BAREMO", E, 970, 400, 180, 48))
    b.append(mx_cell("RUTA", "RUTA", E, 1220, 400, 130, 48))
    for i, (s, t) in enumerate([("TARIFA", "d_tv"), ("d_tv", "VIGENCIA"), ("VIGENCIA", "d_def"),
                                ("d_def", "BAREMO"), ("BAREMO", "d_des"), ("d_des", "DETALLE")]):
        b.append(mx_edge(f"et{i}", s, t))
    b.append(mx_cell("d_posee", "posee", D, 40, 300, 100, 50))
    b.append(mx_cell("d_usa", "usa", D, 560, 300, 90, 50))
    b.append(mx_cell("d_cubre", "cubre", D, 350, 500, 100, 50))
    b.append(mx_cell("d_aplica", "aplica", D, 900, 500, 100, 50))
    b.append(mx_cell("d_ubica", "ubica", D, 1220, 300, 100, 50))
    b += [mx_edge("ep1", "CLIENTE", "d_posee"), mx_edge("ep2", "d_posee", "TARIFA"),
          mx_edge("eu1", "PRODUCTO", "d_usa"), mx_edge("eu2", "d_usa", "BAREMO"),
          mx_edge("eu3", "d_usa", "RUTA"), mx_edge("ec1", "VIGENCIA", "d_cubre"),
          mx_edge("ec2", "d_cubre", "RUTA"), mx_edge("ea1", "BAREMO", "d_aplica"),
          mx_edge("ea2", "d_aplica", "RUTA"), mx_edge("ub1", "CP", "d_ubica"),
          mx_edge("ub2", "d_ubica", "RUTA")]

    b.append(mx_cell("SOLICITUD", "SOLICITUD", EP, 80, 650, 160, 48))
    b.append(mx_cell("d_gen", "genera", D, 280, 646, 100, 54))
    b.append(mx_cell("ARCHIVO", "ARCHIVO", EP, 420, 650, 160, 48))
    b.append(mx_cell("d_val", "se valida", D, 620, 646, 110, 54))
    b.append(mx_cell("VALIDACION", "VALIDACIÓN", EP, 770, 650, 170, 48))
    b += [mx_edge("g1", "SOLICITUD", "d_gen"), mx_edge("g2", "d_gen", "ARCHIVO"),
          mx_edge("v1", "ARCHIVO", "d_val"), mx_edge("v2", "d_val", "VALIDACION")]
    b.append(mx_cell("d_reg", "registra", D, 620, 70, 110, 50))
    b.append(mx_edge("r1", "USUARIO", "d_reg"))
    b.append(mx_edge("r2", "d_reg", "SOLICITUD"))
    b.append(mx_cell("d_ori", "origina", D, 40, 520, 100, 50))
    b.append(mx_edge("o1", "CLIENTE", "d_ori"))
    b.append(mx_edge("o2", "d_ori", "SOLICITUD"))
    b.append(mx_cell("REFORMA", "CONDICIÓN_REFORMA", EP, 500, 800, 240, 48))
    b.append(mx_cell("d_ap", "aplica", D, 560, 740, 110, 50))
    b += [mx_edge("ap1", "CLIENTE", "d_ap"), mx_edge("ap2", "CENTRO", "d_ap"),
          mx_edge("ap3", "PRODUCTO", "d_ap"), mx_edge("ap4", "d_ap", "REFORMA")]
    b.append(mx_cell("leg", "1 / N = cardinalidad. País, moneda y volumétrico NO se copian en RUTA: se consultan en PAÍS y PRODUCTO.", LB, 80, 880, 1300, 28))
    (DIA / "DER_parametrizacion_tarifaria.drawio").write_text(mxfile("DER Chen normalizado", "der-norm", 1500, 980, "".join(b)), encoding="utf-8")
    print("der drawio")


def table_block(prefix, title, rows, x, y, w, fill):
    h = 28 + 20 * len(rows)
    parts = [mx_cell(prefix, title, SW.format(f=fill), x, y, w, h)]
    for i, (text, pk) in enumerate(rows):
        bg = "#FFF8E1" if pk else ("#F4F7F4" if i % 2 == 0 else "#FFFFFF")
        ex = "fontStyle=4;" if pk else ""
        parts.append(
            f'        <mxCell id="{escape(prefix)}-r{i}" value="{escape(text)}" style="{RW.format(bg=bg, ex=ex)}" vertex="1" parent="{escape(prefix)}">\n'
            f'          <mxGeometry y="{28 + i * 20}" width="{w}" height="20" as="geometry" />\n'
            f"        </mxCell>\n"
        )
    return "".join(parts)


def build_mer_drawio():
    g, o, n = "#1B4D3E", "#8A6A2F", "#3D5A80"
    b = [mx_cell("title", "MER / MR normalizado — sin redundancia de llaves", TT, 40, 10, 1400, 32)]
    specs = [
        ("USUARIO", 40, 60, g, 210, [("PK  id_usuario", True), ("nombres", False), ("apellidos", False), ("correo UK", False), ("rol", False), ("estado", False)]),
        ("CLIENTE", 40, 250, g, 210, [("PK  id_cliente", True), ("agrupador UK", False), ("nit_oracle", False), ("nombre_cliente", False), ("tipo_cliente", False), ("ciudad", False), ("direccion", False), ("ejecutivo", False), ("fecha_reforma", False), ("estado", False)]),
        ("CENTRO", 40, 560, g, 210, [("PK  id_centro", True), ("FK  id_cliente", False), ("codigo_centro", False), ("nombre_centro", False)]),
        ("PRODUCTO", 280, 60, g, 220, [("PK  id_producto", True), ("codigo_producto UK", False), ("nombre", False), ("tipo_producto", False), ("tipo_cliente", False), ("tipo_baremo", False), ("factor_volumetrico", False), ("moneda", False)]),
        ("PAIS", 280, 320, g, 220, [("PK  codigo_pais", True), ("nombre_pais", False)]),
        ("CP", 280, 430, g, 220, [("PK  id_codigo_postal", True), ("FK  codigo_pais", False), ("codigo_postal UK", False), ("ciudad", False)]),
        ("TARIFA", 530, 60, g, 210, [("PK  id_tarifa", True), ("FK  id_cliente", False), ("codigo_tarifa", False), ("tipo_tarifa", False), ("estado", False)]),
        ("VIGENCIA", 530, 240, g, 210, [("PK  id_vigencia", True), ("FK  id_tarifa", False), ("fecha_desde", False), ("fecha_hasta", False), ("estado", False)]),
        ("BAREMO", 770, 60, g, 220, [("PK  id_baremo", True), ("FK  id_vigencia", False), ("FK  id_producto", False), ("numero_baremo", False), ("valor_fijo", False), ("valor_minimo", False)]),
        ("DETALLE", 770, 260, g, 220, [("PK  id_detalle", True), ("FK  id_baremo", False), ("hasta_kg", False), ("importe", False), ("sobrepeso", False)]),
        ("RUTA", 770, 430, g, 220, [("PK  id_ruta", True), ("FK  id_vigencia", False), ("FK  id_producto", False), ("FK  id_baremo", False), ("FK  id_cp_origen", False), ("FK  id_cp_destino", False)]),
        ("SOLICITUD", 1030, 60, o, 230, [("PK  id_solicitud", True), ("FK  id_cliente", False), ("FK  id_usuario", False), ("tipo_solicitud", False), ("canal", False), ("fecha_solicitud", False), ("estado", False), ("observaciones", False)]),
        ("ARCHIVO", 1030, 310, o, 230, [("PK  id_archivo", True), ("FK  id_solicitud", False), ("tipo INSUMO|CSV", False), ("nombre_archivo", False), ("ruta", False), ("hash_integridad", False), ("tipo_fichero", False), ("separador", False), ("FK id_usuario_carga", False), ("FK id_usuario_aprueba", False), ("fecha_cargue", False), ("medio_cargue", False), ("resultado_cargue", False), ("registro_sox", False)]),
        ("VALIDACION", 1290, 60, o, 230, [("PK  id_validacion", True), ("FK  id_archivo", False), ("nombre_campo", False), ("fila", False), ("resultado", False), ("detalle", False)]),
        ("REFORMA", 1290, 270, n, 250, [("PK  id_condicion", True), ("tipo DESCUENTO|MINIMA|CARGO", False), ("FK  id_cliente", False), ("FK  id_centro", False), ("FK  id_producto", False), ("codigo_crr", False), ("codigo_zona", False), ("FK  id_cp_origen", False), ("FK  id_cp_destino", False), ("fecha_desde", False), ("fecha_hasta", False), ("fecha_baja", False), ("concepto_facturable", False), ("desde_kg / hasta_kg", False), ("descuento", False), ("rango_kg / minimo", False), ("tasa_manejo", False), ("numero_baremo", False)]),
    ]
    for name, x, y, fill, w, rows in specs:
        b.append(table_block(name, name, rows, x, y, w, fill))
    for i, (s, t) in enumerate([
        ("USUARIO", "SOLICITUD"), ("CLIENTE", "CENTRO"), ("CLIENTE", "TARIFA"), ("CLIENTE", "SOLICITUD"),
        ("TARIFA", "VIGENCIA"), ("VIGENCIA", "BAREMO"), ("VIGENCIA", "RUTA"), ("PRODUCTO", "BAREMO"),
        ("PRODUCTO", "RUTA"), ("BAREMO", "DETALLE"), ("BAREMO", "RUTA"), ("PAIS", "CP"), ("CP", "RUTA"),
        ("SOLICITUD", "ARCHIVO"), ("ARCHIVO", "VALIDACION"), ("CLIENTE", "REFORMA"), ("CENTRO", "REFORMA"),
        ("PRODUCTO", "REFORMA"), ("USUARIO", "ARCHIVO"),
    ]):
        b.append(mx_edge(f"L{i}", s, t, RE))
    (DIA / "MER_modelo_relacional.drawio").write_text(mxfile("MER MR normalizado", "mer-norm", 1600, 1100, "".join(b)), encoding="utf-8")
    print("mer drawio")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
INK_C = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
HDR = "1B4D3E"


def font(run, size=12, bold=False, italic=False, color=INK_C):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def pfmt(p, align="justify", first=True, before=0, after=0):
    f = p.paragraph_format
    f.line_spacing = 2.0
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    f.first_line_indent = Cm(1.27) if first else Cm(0)


def body(doc, text, first=True):
    p = doc.add_paragraph()
    pfmt(p, first=first)
    r = p.add_run(text)
    font(r)


def center(doc, text, size=12, bold=False, italic=False):
    p = doc.add_paragraph()
    pfmt(p, "center", False)
    r = p.add_run(text)
    font(r, size, bold, italic)


def h1(doc, text):
    p = doc.add_paragraph()
    pfmt(p, "center", False, 12)
    r = p.add_run(text)
    font(r, bold=True)


def h2(doc, text):
    p = doc.add_paragraph()
    pfmt(p, "left", False, 10)
    r = p.add_run(text)
    font(r, bold=True)


def cap(doc, kind, n, title):
    p = doc.add_paragraph()
    pfmt(p, "left", False, 10)
    r = p.add_run(f"{kind} {n}")
    font(r, bold=(kind == "Tabla"), italic=(kind == "Figura"))
    p2 = doc.add_paragraph()
    pfmt(p2, "left", False)
    r2 = p2.add_run(title)
    font(r2, italic=True)


def note(doc, text):
    p = doc.add_paragraph()
    pfmt(p, "left", False, 2, 6)
    r = p.add_run("Nota. ")
    font(r, 10, italic=True)
    r2 = p.add_run(text)
    font(r2, 10)


def shade(cell, hx):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hx)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def border(cell):
    b = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8A8A8A")
        b.append(el)
    cell._tc.get_or_add_tcPr().append(b)


def ctext(cell, text, bold=False, size=10, color=INK_C, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    font(r, size, bold, color=color)
    if fill:
        shade(cell, fill)
    border(cell)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        ctext(t.rows[0].cells[i], h, True, 10, WHITE, HDR)
    for ri, row in enumerate(rows):
        fill = "F4F7F4" if ri % 2 == 0 else "FFFFFF"
        for ci, v in enumerate(row):
            ctext(t.rows[ri + 1].cells[ci], str(v), fill=fill)
    return t


def add_page_number(paragraph):
    run = paragraph.add_run()
    a = OxmlElement("w:fldChar")
    a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = " PAGE "
    c = OxmlElement("w:fldChar")
    c.set(qn("w:fldCharType"), "end")
    run._r.extend((a, i, c))
    font(run)


def build_word():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(hp)

    center(doc, "Universidad Militar Nueva Granada", 14, True)
    center(doc, "Facultad de Ingeniería", 12, True)
    center(doc, "Programa de Ingeniería Informática")
    center(doc, "")
    center(doc, "Automatización del proceso de generación y cargue de tarifas al sistema de aplicación de tarifas mediante base de datos e IA", 14, True)
    center(doc, "")
    center(doc, "Actividad 2. Modelo de datos normalizado", italic=True)
    center(doc, "Entidades, atributos, llaves, DER y modelo relacional (sin redundancia)")
    center(doc, "")
    center(doc, "Gina Paola Susatama Molineros", True)
    center(doc, "Jaime Alberto Carreño Camacho", True)
    center(doc, "Bases de Datos I — Osiris Torres Gutiérrez")
    center(doc, "25 de agosto de 2026")
    doc.add_page_break()

    h1(doc, "Criterio de reducción: evitar redundancia")
    body(doc, "Se partió de un modelo de 23 entidades y se normalizó a 15. La regla fue: un hecho de negocio se guarda una sola vez; lo que se pueda consultar con una FK no se copia; lo que sea constante o calculable no se persiste.")
    cap(doc, "Tabla", 1, "Qué se eliminó para no repetir información")
    add_table(doc, ["Se quitó", "Quedó en", "Por qué era redundante"], [
        ["Entidad ROL", "Atributo rol en USUARIO", "Un usuario tiene un solo rol; no hay datos propios del rol."],
        ["ZONA_TARIFARIA y PUNTO_VENTA", "codigo_zona y codigo_crr en CONDICIÓN_REFORMA", "Solo aparecen como códigos opcionales del fichero de reforma."],
        ["ARCHIVO_INSUMO, FICHERO_CSV y CARGUE", "Una entidad ARCHIVO (tipo INSUMO o CSV)", "Los tres describían el mismo objeto archivo y repetían solicitud, nombre y fecha."],
        ["REGLA_VALIDACIÓN y RESULTADO_VALIDACIÓN", "VALIDACIÓN", "El catálogo de reglas se expresaba otra vez en cada hallazgo."],
        ["DESCUENTO, MÍNIMA y CARGO (3 tablas)", "CONDICIÓN_REFORMA con atributo tipo", "Las tres repetían cliente, centro, producto y fechas."],
        ["pais_origen y pais_destino en RUTA", "CÓDIGO_POSTAL.codigo_pais", "El país ya está en el maestro postal."],
        ["conversion_vol y moneda en RUTA", "PRODUCTO.factor_volumetrico y moneda", "Los determina el producto, no la ruta."],
        ["tipo_baremo en BAREMO", "PRODUCTO.tipo_baremo", "KGBU o KILO depende del producto."],
        ["llave_porte", "No se almacena", "Se calcula concatenando rangos al generar el CSV."],
        ["cliente_reforma", "CENTRO (agrupador + código)", "Es la concatenación de dos FKs ya existentes."],
        ["naturaleza, cop_unidad, fracción, %", "No se almacena", "En el instructivo siempre valen CLIENTE o 0."],
    ])
    note(doc, "Se conservan DETALLE_BAREMO y VIGENCIA porque unirlas al padre repetiría el encabezado en cada fila (peor redundancia).")

    h1(doc, "1. Identifique las entidades y sus atributos")
    body(doc, "Quedan 15 entidades fuertes. Cada atributo describe un hecho que no se obtiene por join.")
    cap(doc, "Tabla", 2, "Entidades y atributos del modelo normalizado")
    add_table(doc, ["Entidad", "Atributos"], [
        ["USUARIO", "id_usuario, nombres, apellidos, correo, rol, estado"],
        ["CLIENTE", "id_cliente, agrupador, nit_oracle, nombre_cliente, tipo_cliente, ciudad, direccion, ejecutivo, fecha_reforma, estado"],
        ["CENTRO", "id_centro, id_cliente, codigo_centro, nombre_centro"],
        ["PRODUCTO", "id_producto, codigo_producto, nombre_producto, tipo_producto, tipo_cliente, tipo_baremo, factor_volumetrico, moneda"],
        ["PAÍS", "codigo_pais, nombre_pais"],
        ["CÓDIGO_POSTAL", "id_codigo_postal, codigo_pais, codigo_postal, ciudad"],
        ["TARIFA", "id_tarifa, id_cliente, codigo_tarifa, tipo_tarifa, estado"],
        ["VIGENCIA", "id_vigencia, id_tarifa, fecha_desde, fecha_hasta, estado"],
        ["BAREMO", "id_baremo, id_vigencia, id_producto, numero_baremo, valor_fijo, valor_minimo"],
        ["DETALLE_BAREMO", "id_detalle, id_baremo, hasta_kg, importe, sobrepeso"],
        ["RUTA", "id_ruta, id_vigencia, id_producto, id_baremo, id_cp_origen, id_cp_destino"],
        ["SOLICITUD", "id_solicitud, id_cliente, id_usuario, tipo_solicitud, canal, fecha_solicitud, estado, observaciones"],
        ["ARCHIVO", "id_archivo, id_solicitud, tipo (INSUMO|CSV), nombre_archivo, ruta, hash_integridad, tipo_fichero, separador, id_usuario_carga, id_usuario_aprueba, fecha_cargue, medio_cargue, resultado_cargue, registro_sox"],
        ["VALIDACIÓN", "id_validacion, id_archivo, nombre_campo, fila, resultado, detalle"],
        ["CONDICIÓN_REFORMA", "id_condicion, tipo (DESCUENTO|MINIMA|CARGO), id_cliente, id_centro, id_producto, codigo_crr, codigo_zona, id_cp_origen, id_cp_destino, fecha_desde, fecha_hasta, fecha_baja, concepto_facturable, desde_kg, hasta_kg, descuento, rango_kg, minimo, tasa_manejo, numero_baremo"],
    ])
    note(doc, "En CONDICIÓN_REFORMA, según el tipo se usan descuento, o rango_kg/minimo, o tasa_manejo. Los demás campos específicos quedan nulos. Eso evita tres tablas con las mismas FKs.")

    h1(doc, "2. Determine las llaves primarias y foráneas")
    cap(doc, "Tabla", 3, "Llaves del modelo normalizado")
    add_table(doc, ["Entidad", "PK", "FK"], [
        ["USUARIO", "id_usuario", "—"],
        ["CLIENTE", "id_cliente", "—"],
        ["CENTRO", "id_centro", "id_cliente → CLIENTE"],
        ["PRODUCTO", "id_producto", "—"],
        ["PAÍS", "codigo_pais", "—"],
        ["CÓDIGO_POSTAL", "id_codigo_postal", "codigo_pais → PAÍS"],
        ["TARIFA", "id_tarifa", "id_cliente → CLIENTE (nulo si tarifa general B2C)"],
        ["VIGENCIA", "id_vigencia", "id_tarifa → TARIFA"],
        ["BAREMO", "id_baremo", "id_vigencia → VIGENCIA; id_producto → PRODUCTO"],
        ["DETALLE_BAREMO", "id_detalle", "id_baremo → BAREMO"],
        ["RUTA", "id_ruta", "id_vigencia → VIGENCIA; id_producto → PRODUCTO; id_baremo → BAREMO; id_cp_origen → CÓDIGO_POSTAL; id_cp_destino → CÓDIGO_POSTAL"],
        ["SOLICITUD", "id_solicitud", "id_cliente → CLIENTE; id_usuario → USUARIO"],
        ["ARCHIVO", "id_archivo", "id_solicitud → SOLICITUD; id_usuario_carga → USUARIO; id_usuario_aprueba → USUARIO"],
        ["VALIDACIÓN", "id_validacion", "id_archivo → ARCHIVO"],
        ["CONDICIÓN_REFORMA", "id_condicion", "id_cliente → CLIENTE; id_centro → CENTRO; id_producto → PRODUCTO; id_cp_origen → CÓDIGO_POSTAL; id_cp_destino → CÓDIGO_POSTAL"],
    ])
    note(doc, "UK: CLIENTE.agrupador, PRODUCTO.codigo_producto, USUARIO.correo, CÓDIGO_POSTAL (pais + código). RUTA es única por vigencia + producto + origen + destino.")

    h1(doc, "3. Elabore el diagrama entidad-relación (DER)")
    body(doc, "Notación de Chen. Archivo editable: docs/diagramas/DER_parametrizacion_tarifaria.drawio.")
    p = doc.add_paragraph()
    pfmt(p, "center", False)
    p.add_run().add_picture(str(FIG / "diagrama_er_chen.png"), width=Inches(6.5))
    cap(doc, "Figura", 1, "DER normalizado (15 entidades). País y moneda no se duplican en RUTA.")

    h1(doc, "4. Construya el modelo relacional (MR)")
    body(doc, "Cada relación está en 3FN: no hay atributos transitivos (país vía código postal, moneda vía producto). Archivo editable: docs/diagramas/MER_modelo_relacional.drawio.")
    cap(doc, "Tabla", 4, "Esquema relacional")
    add_table(doc, ["Relación", "Esquema"], [
        ["USUARIO", "USUARIO (id_usuario, nombres, apellidos, correo, rol, estado)"],
        ["CLIENTE", "CLIENTE (id_cliente, agrupador, nit_oracle, nombre_cliente, tipo_cliente, ciudad, direccion, ejecutivo, fecha_reforma, estado)"],
        ["CENTRO", "CENTRO (id_centro, id_cliente, codigo_centro, nombre_centro)"],
        ["PRODUCTO", "PRODUCTO (id_producto, codigo_producto, nombre_producto, tipo_producto, tipo_cliente, tipo_baremo, factor_volumetrico, moneda)"],
        ["PAÍS", "PAÍS (codigo_pais, nombre_pais)"],
        ["CÓDIGO_POSTAL", "CÓDIGO_POSTAL (id_codigo_postal, codigo_pais, codigo_postal, ciudad)"],
        ["TARIFA", "TARIFA (id_tarifa, id_cliente, codigo_tarifa, tipo_tarifa, estado)"],
        ["VIGENCIA", "VIGENCIA (id_vigencia, id_tarifa, fecha_desde, fecha_hasta, estado)"],
        ["BAREMO", "BAREMO (id_baremo, id_vigencia, id_producto, numero_baremo, valor_fijo, valor_minimo)"],
        ["DETALLE_BAREMO", "DETALLE_BAREMO (id_detalle, id_baremo, hasta_kg, importe, sobrepeso)"],
        ["RUTA", "RUTA (id_ruta, id_vigencia, id_producto, id_baremo, id_cp_origen, id_cp_destino)"],
        ["SOLICITUD", "SOLICITUD (id_solicitud, id_cliente, id_usuario, tipo_solicitud, canal, fecha_solicitud, estado, observaciones)"],
        ["ARCHIVO", "ARCHIVO (id_archivo, id_solicitud, tipo, nombre_archivo, ruta, hash_integridad, tipo_fichero, separador, id_usuario_carga, id_usuario_aprueba, fecha_cargue, medio_cargue, resultado_cargue, registro_sox)"],
        ["VALIDACIÓN", "VALIDACIÓN (id_validacion, id_archivo, nombre_campo, fila, resultado, detalle)"],
        ["CONDICIÓN_REFORMA", "CONDICIÓN_REFORMA (id_condicion, tipo, id_cliente, id_centro, id_producto, codigo_crr, codigo_zona, id_cp_origen, id_cp_destino, fecha_desde, fecha_hasta, fecha_baja, concepto_facturable, desde_kg, hasta_kg, descuento, rango_kg, minimo, tasa_manejo, numero_baremo)"],
    ])
    p = doc.add_paragraph()
    pfmt(p, "center", False)
    p.add_run().add_picture(str(FIG / "diagrama_er_crowsfoot.png"), width=Inches(6.5))
    cap(doc, "Figura", 2, "MR normalizado. La flecha apunta al lado N.")

    h2(doc, "Integridad que se mantiene sin copiar datos")
    body(doc, "Ruta única por vigencia, producto y par origen-destino. Todo CSV de rutas exige un ARCHIVO hermano de baremos en la misma solicitud. Producto ya existente en una vigencia implica vigencia nueva. fecha_hasta ≥ fecha_desde. id_usuario_carga distinto de id_usuario_aprueba.")

    out = ROOT / "Actividad2_Entidades_Llaves_DER_MR.docx"
    doc.save(out)
    print("word", out)


if __name__ == "__main__":
    draw_chen()
    draw_crow()
    build_der_drawio()
    build_mer_drawio()
    build_word()
