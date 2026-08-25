#!/usr/bin/env python3
"""Genera el documento Word APA de la Actividad Integradora (entrega 2)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Emu

OUT = Path("/workspace/docs/Actividad_Integradora_UMNG_Entidades_Atributos_ER.docx")
FIG_CHEN = Path("/workspace/docs/figuras/diagrama_er_chen.png")
FIG_CROW = Path("/workspace/docs/figuras/diagrama_er_crowsfoot.png")

GREEN = RGBColor(0x1B, 0x4D, 0x3E)
INK = RGBColor(0x00, 0x00, 0x00)
HEADER_BG = "1B4D3E"


def set_run_font(run, size=12, bold=False, italic=False, color=INK, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_format(p, align="justify", first_line=True, space_before=0, space_after=0, line=2.0, keep=False):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.line_spacing = line
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Cm(1.27) if first_line else Cm(0)
    if keep:
        pf.keep_with_next = True


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, 12)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8A8A8A")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, bold=False, size=10, color=INK, align="left", fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_center(doc, text, size=12, bold=False, italic=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="center", first_line=False, space_before=space_before, space_after=space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def heading1(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="center", first_line=False, space_before=12, space_after=0)
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def heading2(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="left", first_line=False, space_before=12, space_after=0)
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def heading3(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="left", first_line=False, space_before=8, space_after=0)
    run = p.add_run(text)
    set_run_font(run, bold=True, italic=True)
    return p


def caption_table(doc, number, title):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="left", first_line=False, space_before=12, space_after=0)
    r1 = p.add_run(f"Tabla {number}")
    set_run_font(r1, bold=True)
    p2 = doc.add_paragraph()
    set_paragraph_format(p2, align="left", first_line=False, space_before=0, space_after=6)
    r2 = p2.add_run(title)
    set_run_font(r2, italic=True)


def caption_figure(doc, number, title):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="left", first_line=False, space_before=6, space_after=0)
    r1 = p.add_run(f"Figura {number}")
    set_run_font(r1, italic=True)
    p2 = doc.add_paragraph()
    set_paragraph_format(p2, align="left", first_line=False, space_before=0, space_after=12)
    r2 = p2.add_run(title)
    set_run_font(r2)


def note_table(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align="left", first_line=False, space_before=2, space_after=8)
    r = p.add_run("Nota. ")
    set_run_font(r, size=10, italic=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True, size=10, color=RGBColor(255, 255, 255), align="center", fill=HEADER_BG)
    for r_i, row in enumerate(rows):
        fill = "F4F7F4" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            align = "left" if c_i > 0 else "left"
            cell_text(table.rows[r_i + 1].cells[c_i], str(val), size=10, fill=fill, align=align)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    return table


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing = 2.0
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    add_page_number(hp)


def build():
    doc = Document()
    configure_styles(doc)

    # ===================== PORTADA =====================
    for _ in range(2):
        add_center(doc, "")
    add_center(doc, "Universidad Militar Nueva Granada", size=14, bold=True)
    add_center(doc, "Facultad de Ingeniería", size=12, bold=True)
    add_center(doc, "Programa de Ingeniería Informática", size=12)
    add_center(doc, "")
    add_center(
        doc,
        "Automatización del proceso de generación y cargue de tarifas al sistema de aplicación de tarifas mediante base de datos e IA",
        size=14,
        bold=True,
    )
    add_center(doc, "")
    add_center(doc, "Proyecto de Actividad Integradora", size=12, italic=True)
    add_center(doc, "Entrega 2. Identificación de entidades, atributos, relaciones y diagrama entidad-relación", size=12)
    add_center(doc, "")
    add_center(doc, "Gina Paola Susatama Molineros", bold=True)
    add_center(doc, "est.gina.susatama@unimilitar.edu.co", size=12, italic=True)
    add_center(doc, "Jaime Alberto Carreño Camacho", bold=True)
    add_center(doc, "est.jaime.carreno@unimilitar.edu.co", size=12, italic=True)
    add_center(doc, "")
    add_center(doc, "Asignatura: Bases de Datos I")
    add_center(doc, "Docente: Osiris Torres Gutiérrez")
    add_center(doc, "")
    add_center(doc, "Bogotá, D. C., Colombia")
    add_center(doc, "25 de agosto de 2026")

    doc.add_page_break()

    heading1(doc, "Tabla de Contenido")
    toc_items = [
        "Resumen",
        "Introducción",
        "Identificación del Proyecto de Actividad Integradora",
        "    Integrantes del grupo de trabajo",
        "    Nombre del proyecto",
        "    Planteamiento del problema",
        "    Objetivo general",
        "    Objetivos específicos",
        "    Justificación",
        "    Motivación intrínseca",
        "    Asignaturas involucradas",
        "Método de Identificación del Modelo de Datos",
        "    Mapeo del proceso actual a entidades",
        "Entidades del Proyecto",
        "Atributos del Proyecto",
        "Relaciones del Proyecto",
        "    Reglas de integridad relevantes",
        "Diagrama Entidad-Relación",
        "    Lectura del diagrama",
        "Conclusiones",
        "Referencias",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        set_paragraph_format(p, align="left", first_line=False)
        run = p.add_run(item)
        set_run_font(run)

    doc.add_page_break()

    # ===================== RESUMEN =====================
    heading1(doc, "Resumen")
    add_body(
        doc,
        "El presente documento desarrolla la segunda entrega del Proyecto de Actividad Integradora, "
        "conservando el grupo de trabajo y todos los ítems diligenciados en el formato de la entrega 1, "
        "e incorporando la identificación de entidades, atributos y relaciones del sistema propuesto, "
        "junto con el diagrama entidad-relación. El proyecto busca automatizar el proceso manual de "
        "parametrización tarifaria para el envío de paquetes en el sistema Alertran Padua, en el cual "
        "los insumos llegan en múltiples archivos Excel que deben convertirse a CSV con un formato "
        "específico, validarse contra los lineamientos del sistema y cargarse por FTP o por importación "
        "de maestros. A partir del análisis del proceso actual se diseñó un modelo relacional de 23 "
        "entidades, agrupadas en maestros de negocio, estructura tarifaria, flujo de automatización y "
        "reforma tarifaria. El modelo centraliza la información, garantiza integridad referencial, "
        "registra quién aprueba y quién carga cada tarifa, y deja trazabilidad de archivos origen, "
        "ficheros generados, validaciones y resultado del cargue.",
    )
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False)
    r = p.add_run("Palabras clave: ")
    set_run_font(r, italic=True)
    r2 = p.add_run("modelo entidad-relación, parametrización tarifaria, base de datos, automatización, CSV, integridad de datos.")
    set_run_font(r2)

    # ===================== INTRODUCCIÓN =====================
    heading1(doc, "Introducción")
    add_body(
        doc,
        "Una vez diligenciado el formato de proyecto de la Actividad Integradora en la entrega 1, "
        "esta segunda entrega mantiene el mismo grupo de trabajo y transcribe la totalidad de los "
        "ítems del formulario, para luego identificar las entidades, los atributos y las relaciones "
        "del dominio propuesto. El producto se presenta en un documento Word elaborado según las "
        "normas APA (7.ª edición), e incluye el diagrama entidad-relación del sistema.",
    )
    add_body(
        doc,
        "El dominio de negocio es el proceso tarifario de envío de paquetes. En la operación actual, "
        "la parametrización de tarifas en Alertran Padua se realiza de forma predominantemente manual: "
        "se reciben varios Excel, se construyen llaves de porte, se asignan baremos, se arman ficheros "
        "de rutas y baremos (o de reforma: descuentos, mínima de despacho y cargo por manejo) y se "
        "convierten a CSV con separador de listas tipo tubería, para luego cargarlos al sistema "
        "respetando vigencias, productos, códigos postales y controles SOX (Deprisa, 2026). La dolencia "
        "que el proyecto busca resolver es precisamente esa conversión y ese cargue manual, propensos "
        "a error humano y dependientes del conocimiento tácito del analista.",
    )
    add_body(
        doc,
        "El documento se organiza en cuatro apartados. Primero se presenta la identificación del "
        "proyecto con todos los campos del formato. Segundo se describe el método de identificación "
        "del modelo de datos. Tercero se detallan entidades, atributos, relaciones, cardinalidades y "
        "los diagramas entidad-relación. Cuarto se cierran conclusiones y referencias.",
    )

    # ===================== IDENTIFICACIÓN DEL PROYECTO =====================
    heading1(doc, "Identificación del Proyecto de Actividad Integradora")
    add_body(
        doc,
        "Esta sección reproduce, sin alterar el sentido original, los ítems diligenciados en el "
        "Formato de Proyecto de Actividad Integradora de la Universidad Militar Nueva Granada "
        "(entrega 1). Se conserva el grupo de trabajo registrado en dicho formato.",
        first_line=True,
    )

    heading2(doc, "Integrantes del grupo de trabajo")
    add_body(
        doc,
        "El grupo de trabajo se mantiene respecto de la entrega 1. Los integrantes son los siguientes:",
    )
    caption_table(doc, 1, "Integrantes del grupo de trabajo")
    add_table(
        doc,
        ["N.º", "Nombre y apellido", "Correo institucional"],
        [
            ["1", "Gina Paola Susatama Molineros", "est.gina.susatama@unimilitar.edu.co"],
            ["2", "Jaime Alberto Carreño Camacho", "est.jaime.carreno@unimilitar.edu.co"],
        ],
        col_widths=[1.5, 7.5, 7.0],
    )
    note_table(doc, "Elaboración propia a partir del formato de la entrega 1.")

    heading2(doc, "Nombre del proyecto")
    add_body(
        doc,
        "Automatización del proceso de generación y cargue de tarifas al sistema de aplicación de tarifas mediante base de datos e IA.",
    )

    heading2(doc, "Planteamiento del problema")
    add_body(
        doc,
        "Actualmente, la parametrización de tarifas en el sistema de aplicación tarifaria depende de "
        "un proceso manual: se reciben varios archivos Excel como insumos tarifarios de distintas "
        "fuentes, cada uno debe convertirse manualmente a formato CSV siguiendo reglas específicas "
        "del sistema (columnas, orden, tipos de dato, codificación), y luego cargarse de forma "
        "individual. Este proceso es dispendioso, propenso a errores humanos, ya que la construcción "
        "de los ficheros puede alterar la información y afectar su integridad, y depende del "
        "conocimiento tácito de la persona que lo ejecuta, lo que genera reprocesos y riesgo "
        "operativo cuando dicha persona no está disponible.",
    )
    add_body(
        doc,
        "El instructivo interno de parametrización tarifaria confirma esa complejidad: existen tarifas "
        "generales, especiales y de reforma; clientes B2B y B2C; productos nacional, internacional y "
        "carga; ficheros de rutas y baremos que deben cargarse de manera simultánea en el FTP; y "
        "ficheros de reforma (descuentos, mínima de despacho y cargo por manejo) con validaciones de "
        "longitud, formato de fecha, separador decimal y existencia de maestros (Deprisa, 2026). "
        "La condición que debe mejorarse es, por tanto, el ciclo completo de preparación, validación "
        "y cargue de esos ficheros, no un ajuste puntual de una sola tarifa.",
    )

    heading2(doc, "Objetivo general")
    add_body(
        doc,
        "Diseñar e implementar una base de datos y un flujo automatizado, apoyado en programación e "
        "inteligencia artificial, que estandarice la conversión de los archivos Excel de tarifas al "
        "formato CSV requerido por el sistema, facilite su cargue y garantice la integridad y "
        "trazabilidad de la información, reduciendo los errores manuales y el tiempo de procesamiento "
        "del proceso actual.",
    )

    heading2(doc, "Objetivos específicos")
    add_body(
        doc,
        "1. Diseñar el modelo de base de datos relacional que centralice, valide y estandarice la "
        "información tarifaria proveniente de los distintos archivos Excel de origen.",
    )
    add_body(
        doc,
        "2. Desarrollar un módulo de conversión asistido por IA que transforme los archivos Excel en "
        "el formato CSV exigido por el sistema, validando de manera automática las reglas de formato "
        "(columnas, tipos de dato, orden y codificación).",
    )
    add_body(
        doc,
        "3. Implementar un programa de cargue que tome el archivo CSV generado, verifique su "
        "cumplimiento con los lineamientos del sistema antes de subirlo y registre en la base de "
        "datos la trazabilidad de cada cargue, incluyendo quién aprueba cada tarifa, quién la carga, "
        "fecha, archivo origen y resultado.",
    )

    heading2(doc, "Justificación")
    add_body(
        doc,
        "El proyecto responde a una necesidad real identificada en el rol del colaborador que elabora "
        "y ejecuta el proceso, donde el cargue de tarifas se realiza hoy de forma manual a partir de "
        "múltiples archivos Excel. Automatizar este flujo reduce el tiempo dedicado a tareas "
        "repetitivas, disminuye el riesgo de errores humanos en la conversión y el cargue, y deja "
        "trazabilidad de la información mediante una base de datos, lo cual mejora la calidad del "
        "proceso y libera tiempo para actividades de mayor valor. Adicionalmente, permite aplicar de "
        "manera integrada los conocimientos de bases de datos, programación, seguridad informática y "
        "análisis de sistemas de información desarrollados en la carrera.",
    )
    add_body(
        doc,
        "Desde el punto de vista de integridad de datos, el modelo relacional es el mecanismo adecuado "
        "para impedir rutas duplicadas, vigencias inconsistentes, productos inexistentes o ficheros "
        "CSV que no respeten longitudes y formatos. Desde el punto de vista de seguridad, el registro "
        "de aprobación y de ejecución del cargue atiende el control SOX AVH-263-F y el principio de "
        "separación de funciones (Deprisa, 2026; Pressman y Maxim, 2020).",
    )

    heading2(doc, "Motivación intrínseca")
    add_body(
        doc,
        "Vivo en primera persona la ineficiencia de este proceso manual en mi trabajo diario y quiero "
        "aplicar lo aprendido en la carrera para resolver un problema concreto de mi entorno laboral. "
        "Me motiva fortalecer mis habilidades en bases de datos y programación, aportar una solución "
        "que pueda ser útil real y directamente en la compañía, y demostrar cómo la ingeniería "
        "informática puede mejorar procesos existentes.",
    )

    heading2(doc, "Asignaturas involucradas")
    caption_table(doc, 2, "Relación de asignaturas que se involucran en el proyecto")
    add_table(
        doc,
        ["Asignatura", "Aporte al proyecto", "Docente"],
        [
            [
                "Bases de Datos I",
                "Aporta el diseño del modelo entidad-relación y la implementación de la base de datos que centraliza y valida la información tarifaria.",
                "Osiris Torres Gutiérrez",
            ],
            [
                "Programación III",
                "Aporta las bases para desarrollar el módulo de conversión de Excel a CSV y el programa de cargue automatizado al sistema.",
                "Monica Barrios",
            ],
            [
                "Análisis y Diseño de Sistemas de Información",
                "Aporta la metodología para levantar requisitos, modelar el proceso actual frente al propuesto y diseñar la solución de software.",
                "Liliana Constanza Contreras Callejas",
            ],
            [
                "Seguridad informática",
                "Aporta los principios de control de acceso, protección e integridad de la información aplicados al flujo de conversión y cargue de tarifas, garantizando que solo usuarios autorizados aprueben y carguen tarifas en el sistema.",
                "Sin asignar",
            ],
        ],
        col_widths=[4.0, 8.5, 4.5],
    )
    note_table(doc, "Elaboración propia a partir del formato de la entrega 1.")

    # ===================== MÉTODO =====================
    heading1(doc, "Método de Identificación del Modelo de Datos")
    add_body(
        doc,
        "La identificación de entidades, atributos y relaciones se realizó con el enfoque clásico del "
        "modelo entidad-relación (Chen, 1976; Elmasri y Navathe, 2016). Se partió de los sustantivos "
        "del dominio (cliente, producto, tarifa, vigencia, ruta, baremo, fichero, cargue) y de los "
        "artefactos del proceso (Excel de solicitud, hoja tarifaria, aval de optimización, CSV de "
        "rutas, CSV de baremos y ficheros de reforma). Cada candidato se clasificó como entidad si "
        "representa un objeto persistente con identidad propia; como atributo si describe una "
        "propiedad; y como relación si conecta dos o más entidades con una regla de negocio.",
    )
    add_body(
        doc,
        "Se contrastó el modelo con los lineamientos del sistema destino. Para el maestro 1.2.1 se "
        "exigieron vigencia, código de tarifa, producto, país, código postal, número de baremo, "
        "conversión volumétrica y moneda. Para baremos se exigieron tipo (KGBU o KILO), hasta, "
        "importe y sobrepeso. Para reforma tarifaria se exigieron agrupador, centro, producto, "
        "fechas desde/hasta/baja y los campos propios de descuento, mínima y cargo por manejo "
        "(Deprisa, 2026). Las reglas de validación no se dejaron implícitas en código: se modelaron "
        "como entidad, de modo que el módulo de conversión asistido por IA y el programa de cargue "
        "consulten el mismo catálogo de lineamientos.",
    )

    heading2(doc, "Mapeo del proceso actual a entidades")
    caption_table(doc, 3, "Correspondencia entre el proceso tarifario manual y las entidades del modelo")
    add_table(
        doc,
        ["Elemento del proceso actual", "Entidad o conjunto de entidades"],
        [
            ["Ticket Dynamo o correo de solicitud B2C/B2B", "SOLICITUD"],
            ["Excel de insumos, hoja tarifaria, aval y archivo plano", "ARCHIVO_INSUMO"],
            ["Analista que construye y carga; supervisor que aprueba", "USUARIO, ROL"],
            ["Cliente B2B o B2C, NIT/Oracle, agrupador y centro", "CLIENTE, CENTRO"],
            ["Códigos de producto nacional, internacional y carga", "PRODUCTO"],
            ["Tarifa general, especial o reforma y su código", "TARIFA"],
            ["Fechas desde/hasta y no reutilizar vigencia con los mismos productos", "VIGENCIA"],
            ["Malla origen-destino (país y código postal)", "PAÍS, CÓDIGO_POSTAL, RUTA"],
            ["Llave porte, numeración de baremo, hasta, importe y kilo adicional", "BAREMO, DETALLE_BAREMO"],
            ["CSV de rutas y baremos (separador |, sin encabezados)", "FICHERO_CSV"],
            ["Lineamientos de longitud, tipo, obligatorio y formato", "REGLA_VALIDACIÓN, RESULTADO_VALIDACIÓN"],
            ["Cargue FTP simultáneo o importación 12.5 y log de error", "CARGUE"],
            ["Control SOX AVH-263-F", "CARGUE (registro_sox)"],
            ["Descuentos, mínima de despacho y cargo por manejo", "DESCUENTO_TARIFA, MÍNIMA_DESPACHO, CARGO_MANEJO"],
            ["Zona tarifaria y punto de venta (CRR)", "ZONA_TARIFARIA, PUNTO_VENTA"],
        ],
        col_widths=[8.0, 8.5],
    )
    note_table(doc, "Elaboración propia con base en Deprisa (2026) y en el formato de la entrega 1.")

    # ===================== ENTIDADES =====================
    heading1(doc, "Entidades del Proyecto")
    add_body(
        doc,
        "Se identificaron 23 entidades, organizadas en cuatro grupos: seguridad, maestros de negocio, "
        "estructura tarifaria y flujo de automatización (incluye reforma). A continuación se define "
        "cada entidad, su tipo y su justificación en el dominio.",
    )

    heading2(doc, "Entidades de seguridad y actores")
    caption_table(doc, 4, "Entidades de seguridad y actores")
    add_table(
        doc,
        ["Entidad", "Definición", "Tipo"],
        [
            ["ROL", "Perfil de autorización (analista tarifario, aprobador, administrador).", "Fuerte"],
            ["USUARIO", "Persona que registra solicitudes, aprueba tarifas o ejecuta el cargue.", "Fuerte"],
        ],
        col_widths=[3.5, 10.5, 2.5],
    )
    note_table(doc, "Tipo fuerte indica que la entidad tiene llave primaria propia.")

    heading2(doc, "Entidades maestras de negocio")
    caption_table(doc, 5, "Entidades maestras de negocio")
    add_table(
        doc,
        ["Entidad", "Definición", "Tipo"],
        [
            ["CLIENTE", "Persona o empresa B2B/B2C dueña de la negociación tarifaria. Se identifica con agrupador de 8 caracteres y datos Oracle/NIT.", "Fuerte"],
            ["CENTRO", "Centro operativo de 2 caracteres asociado a un agrupador.", "Fuerte"],
            ["PRODUCTO", "Servicio de envío (nacional, internacional o carga) con código de 4 dígitos, factor volumétrico y moneda.", "Fuerte"],
            ["PAÍS", "Catálogo de país con código de 3 caracteres.", "Fuerte"],
            ["CÓDIGO_POSTAL", "Código postal de 6 dígitos o comodín (*) usado como origen o destino de la ruta.", "Fuerte"],
            ["ZONA_TARIFARIA", "Zona de 4 caracteres usada en reforma tarifaria.", "Fuerte"],
            ["PUNTO_VENTA", "Punto de venta o CRR de 8 caracteres, opcional en reforma.", "Fuerte"],
        ],
        col_widths=[3.5, 10.5, 2.5],
    )
    note_table(doc, "Los tamaños de código reproducen las validaciones del sistema destino (Deprisa, 2026).")

    heading2(doc, "Entidades de estructura tarifaria")
    caption_table(doc, 6, "Entidades de la estructura tarifaria (maestro 1.2.1)")
    add_table(
        doc,
        ["Entidad", "Definición", "Tipo"],
        [
            ["TARIFA", "Código tarifario de naturaleza CLIENTE, clasificado como GENERAL, ESPECIAL o REFORMA.", "Fuerte"],
            ["VIGENCIA", "Periodo de aplicación de una tarifa. El sistema no reprocesa productos ya existentes en la misma vigencia.", "Fuerte"],
            ["BAREMO", "Escala de portes numerada (1 a 999) asociada a una vigencia y a un producto.", "Fuerte"],
            ["DETALLE_BAREMO", "Rango de peso (hasta), importe, sobrepeso y campos fijos en 0 cuando no aplican.", "Débil de BAREMO"],
            ["RUTA", "Combinación origen-destino de una vigencia y un producto, ligada a un baremo, conversión volumétrica y moneda.", "Fuerte"],
        ],
        col_widths=[3.5, 10.5, 2.5],
    )
    note_table(
        doc,
        "RUTA y BAREMO son dependientes de negocio: no puede existir un baremo huérfano en el cargue FTP, porque todo baremo debe estar asociado a una ruta (Deprisa, 2026).",
    )

    heading2(doc, "Entidades del flujo de automatización y de reforma")
    caption_table(doc, 7, "Entidades del flujo de conversión, validación, cargue y reforma")
    add_table(
        doc,
        ["Entidad", "Definición", "Tipo"],
        [
            ["SOLICITUD", "Caso de parametrización recibido por correo (B2C) o por Dynamo (B2B).", "Fuerte"],
            ["ARCHIVO_INSUMO", "Excel u otro adjunto de la solicitud (plano, hoja tarifaria, aval).", "Fuerte"],
            ["FICHERO_CSV", "Archivo generado en el formato exigido (rutas, baremos, descuentos, mínima o cargo).", "Fuerte"],
            ["REGLA_VALIDACIÓN", "Lineamiento parametrizable del sistema (longitud, tipo, obligatoriedad, formato).", "Fuerte"],
            ["RESULTADO_VALIDACIÓN", "Hallazgo de una regla sobre una fila de un fichero.", "Fuerte"],
            ["CARGUE", "Intento de subida al FTP o a la importación 12.5, con log, medio y control SOX.", "Fuerte"],
            ["DESCUENTO_TARIFA", "Porcentaje o valor de descuento de reforma por cliente, centro y producto.", "Fuerte"],
            ["MÍNIMA_DESPACHO", "Valor mínimo de despacho de reforma, con rango de kg y baremo opcional.", "Fuerte"],
            ["CARGO_MANEJO", "Tasa e importe mínimo de cargo por manejo de reforma.", "Fuerte"],
        ],
        col_widths=[4.0, 10.0, 2.5],
    )
    note_table(doc, "Elaboración propia.")

    # ===================== ATRIBUTOS =====================
    heading1(doc, "Atributos del Proyecto")
    add_body(
        doc,
        "Los atributos se derivaron de las columnas de los ficheros de cargue y de los datos de "
        "control que el objetivo específico 3 exige persistir (quién aprueba, quién carga, fecha, "
        "archivo origen y resultado). En las tablas siguientes, PK es llave primaria, FK llave "
        "foránea, UK llave única y NN indica obligatoriedad.",
    )

    heading2(doc, "Atributos de seguridad, cliente y producto")
    caption_table(doc, 8, "Atributos de ROL, USUARIO, CLIENTE, CENTRO y PRODUCTO")
    add_table(
        doc,
        ["Entidad", "Atributo", "Tipo / dominio", "Restricción"],
        [
            ["ROL", "id_rol", "Entero", "PK"],
            ["ROL", "nombre_rol", "Texto(40)", "UK, NN"],
            ["ROL", "descripcion", "Texto(200)", "Opcional"],
            ["USUARIO", "id_usuario", "Entero", "PK"],
            ["USUARIO", "id_rol", "Entero", "FK → ROL, NN"],
            ["USUARIO", "nombres", "Texto(80)", "NN"],
            ["USUARIO", "apellidos", "Texto(80)", "NN"],
            ["USUARIO", "correo", "Texto(120)", "UK, NN"],
            ["USUARIO", "estado", "Activo/Inactivo", "NN"],
            ["CLIENTE", "id_cliente", "Entero", "PK"],
            ["CLIENTE", "agrupador", "Texto(8)", "UK, NN"],
            ["CLIENTE", "nit_oracle", "Texto(20)", "UK"],
            ["CLIENTE", "nombre_cliente", "Texto(150)", "NN"],
            ["CLIENTE", "tipo_cliente", "B2B / B2C", "NN"],
            ["CLIENTE", "ciudad", "Texto(80)", "Opcional"],
            ["CLIENTE", "direccion", "Texto(200)", "Opcional"],
            ["CLIENTE", "ejecutivo_cuenta", "Texto(120)", "Opcional"],
            ["CLIENTE", "fecha_reforma", "Fecha", "Opcional"],
            ["CLIENTE", "cliente_reforma", "Texto(10)", "Agrupador+centro"],
            ["CLIENTE", "estado", "Activo/Inactivo", "NN"],
            ["CENTRO", "id_centro", "Entero", "PK"],
            ["CENTRO", "id_cliente", "Entero", "FK → CLIENTE, NN"],
            ["CENTRO", "codigo_centro", "Texto(2)", "NN"],
            ["CENTRO", "nombre_centro", "Texto(80)", "Opcional"],
            ["PRODUCTO", "id_producto", "Entero", "PK"],
            ["PRODUCTO", "codigo_producto", "Texto(4)", "UK, NN"],
            ["PRODUCTO", "nombre_producto", "Texto(120)", "NN"],
            ["PRODUCTO", "tipo_producto", "Nacional / Internacional / Carga", "NN"],
            ["PRODUCTO", "tipo_cliente", "B2B / B2C", "NN"],
            ["PRODUCTO", "tipo_baremo", "KGBU / KILO", "NN"],
            ["PRODUCTO", "factor_volumetrico", "400, 222, 167, 200 o 0", "NN"],
            ["PRODUCTO", "moneda", "COP / USD", "NN"],
        ],
        col_widths=[3.2, 4.2, 5.6, 3.5],
    )
    note_table(
        doc,
        "El factor volumétrico sigue el instructivo: 400 mercancía, 222 premium y bolsa, 167 carga, 200 internacional y 0 clientes especiales. La moneda es COP para nacional y carga, y USD para internacional (Deprisa, 2026).",
    )

    heading2(doc, "Atributos geográficos y de tarifa")
    caption_table(doc, 9, "Atributos de PAÍS, CÓDIGO_POSTAL, ZONA_TARIFARIA, PUNTO_VENTA, TARIFA y VIGENCIA")
    add_table(
        doc,
        ["Entidad", "Atributo", "Tipo / dominio", "Restricción"],
        [
            ["PAÍS", "codigo_pais", "Texto(3)", "PK"],
            ["PAÍS", "nombre_pais", "Texto(80)", "NN"],
            ["CÓDIGO_POSTAL", "id_codigo_postal", "Entero", "PK"],
            ["CÓDIGO_POSTAL", "codigo_pais", "Texto(3)", "FK → PAÍS, NN"],
            ["CÓDIGO_POSTAL", "codigo_postal", "Texto(6) o *", "UK con país, NN"],
            ["CÓDIGO_POSTAL", "ciudad", "Texto(80)", "Opcional"],
            ["ZONA_TARIFARIA", "id_zona", "Entero", "PK"],
            ["ZONA_TARIFARIA", "codigo_zona", "Texto(4)", "UK, NN"],
            ["ZONA_TARIFARIA", "nombre_zona", "Texto(80)", "Opcional"],
            ["PUNTO_VENTA", "id_punto_venta", "Entero", "PK"],
            ["PUNTO_VENTA", "codigo_crr", "Texto(8)", "UK, NN"],
            ["PUNTO_VENTA", "nombre_punto", "Texto(120)", "Opcional"],
            ["TARIFA", "id_tarifa", "Entero", "PK"],
            ["TARIFA", "id_cliente", "Entero", "FK → CLIENTE, nulo en general B2C"],
            ["TARIFA", "codigo_tarifa", "Texto(30)", "NN"],
            ["TARIFA", "tipo_tarifa", "GENERAL / ESPECIAL / REFORMA", "NN"],
            ["TARIFA", "naturaleza", "CLIENTE", "NN"],
            ["TARIFA", "estado", "Activa/Inactiva", "NN"],
            ["VIGENCIA", "id_vigencia", "Entero", "PK"],
            ["VIGENCIA", "id_tarifa", "Entero", "FK → TARIFA, NN"],
            ["VIGENCIA", "fecha_desde", "Fecha dd/mm/aaaa", "NN"],
            ["VIGENCIA", "fecha_hasta", "Fecha", "Opcional; no inferior a desde"],
            ["VIGENCIA", "estado", "Activa/Cerrada", "NN"],
        ],
        col_widths=[3.5, 4.0, 5.5, 3.5],
    )
    note_table(doc, "Elaboración propia.")

    heading2(doc, "Atributos de baremo y ruta")
    caption_table(doc, 10, "Atributos de BAREMO, DETALLE_BAREMO y RUTA")
    add_table(
        doc,
        ["Entidad", "Atributo", "Tipo / dominio", "Restricción"],
        [
            ["BAREMO", "id_baremo", "Entero", "PK"],
            ["BAREMO", "id_vigencia", "Entero", "FK → VIGENCIA, NN"],
            ["BAREMO", "id_producto", "Entero", "FK → PRODUCTO, NN"],
            ["BAREMO", "numero_baremo", "Entero 1–999", "NN; único por vigencia y producto"],
            ["BAREMO", "tipo_baremo", "KGBU / KILO", "NN"],
            ["BAREMO", "valor_fijo", "Decimal", "Default 0"],
            ["BAREMO", "valor_minimo", "Decimal", "Default 0"],
            ["DETALLE_BAREMO", "id_detalle", "Entero", "PK"],
            ["DETALLE_BAREMO", "id_baremo", "Entero", "FK → BAREMO, NN"],
            ["DETALLE_BAREMO", "hasta_kg", "Decimal / 9999999", "NN; 1, 5, 30, 100 o 9999999"],
            ["DETALLE_BAREMO", "importe", "Decimal", "NN; flete del rango"],
            ["DETALLE_BAREMO", "sobrepeso", "Decimal", "Kilo adicional"],
            ["DETALLE_BAREMO", "cop_unidad", "Decimal", "Default 0"],
            ["DETALLE_BAREMO", "fraccion", "Decimal", "Default 0"],
            ["DETALLE_BAREMO", "porcentaje", "Decimal", "Default 0"],
            ["RUTA", "id_ruta", "Entero", "PK"],
            ["RUTA", "id_vigencia", "Entero", "FK → VIGENCIA, NN"],
            ["RUTA", "id_producto", "Entero", "FK → PRODUCTO, NN"],
            ["RUTA", "id_baremo", "Entero", "FK → BAREMO, NN"],
            ["RUTA", "pais_origen", "Texto(3)", "FK → PAÍS, NN"],
            ["RUTA", "id_cp_origen", "Entero", "FK → CÓDIGO_POSTAL"],
            ["RUTA", "pais_destino", "Texto(3)", "FK → PAÍS, NN"],
            ["RUTA", "id_cp_destino", "Entero", "FK → CÓDIGO_POSTAL"],
            ["RUTA", "conversion_vol", "Entero", "NN; copia del factor del producto"],
            ["RUTA", "moneda", "COP / USD", "NN"],
            ["RUTA", "llave_porte", "Texto(120)", "Concatenación de rangos para asignar baremo"],
        ],
        col_widths=[3.5, 4.0, 5.0, 4.0],
    )
    note_table(
        doc,
        "La ruta no puede estar duplicada en una misma vigencia y producto. Delegación origen y destino se modelan como nulos permanentes, porque el instructivo exige celdas en blanco (Deprisa, 2026).",
    )

    heading2(doc, "Atributos del flujo de automatización")
    caption_table(doc, 11, "Atributos de SOLICITUD, ARCHIVO_INSUMO, FICHERO_CSV, REGLA_VALIDACIÓN, RESULTADO_VALIDACIÓN y CARGUE")
    add_table(
        doc,
        ["Entidad", "Atributo", "Tipo / dominio", "Restricción"],
        [
            ["SOLICITUD", "id_solicitud", "Entero", "PK"],
            ["SOLICITUD", "id_cliente", "Entero", "FK → CLIENTE"],
            ["SOLICITUD", "id_usuario_registro", "Entero", "FK → USUARIO, NN"],
            ["SOLICITUD", "tipo_solicitud", "B2C_GENERAL / B2B_ESPECIAL / REFORMA", "NN"],
            ["SOLICITUD", "canal_recepcion", "Correo / Dynamo", "NN"],
            ["SOLICITUD", "fecha_solicitud", "Fecha-hora", "NN"],
            ["SOLICITUD", "estado", "Recibida, validada, rechazada, convertida, cargada", "NN"],
            ["SOLICITUD", "observaciones", "Texto largo", "Opcional"],
            ["ARCHIVO_INSUMO", "id_archivo", "Entero", "PK"],
            ["ARCHIVO_INSUMO", "id_solicitud", "Entero", "FK → SOLICITUD, NN"],
            ["ARCHIVO_INSUMO", "nombre_archivo", "Texto(200)", "NN"],
            ["ARCHIVO_INSUMO", "tipo_insumo", "Excel tarifas / hoja tarifaria / aval / plano", "NN"],
            ["ARCHIVO_INSUMO", "ruta_almacenamiento", "Texto(260)", "NN"],
            ["ARCHIVO_INSUMO", "fecha_carga", "Fecha-hora", "NN"],
            ["ARCHIVO_INSUMO", "hash_integridad", "Texto(64)", "NN; protege alteración"],
            ["FICHERO_CSV", "id_fichero", "Entero", "PK"],
            ["FICHERO_CSV", "id_solicitud", "Entero", "FK → SOLICITUD, NN"],
            ["FICHERO_CSV", "tipo_fichero", "RUTAS / BAREMOS / DESCUENTOS / MINIMA / CARGO / CLIENTE", "NN"],
            ["FICHERO_CSV", "nombre_fichero", "Texto(120)", "NN; patrón ddmmaa + consecutivo"],
            ["FICHERO_CSV", "separador", "Carácter", "NN; valor |"],
            ["FICHERO_CSV", "fecha_generacion", "Fecha-hora", "NN"],
            ["FICHERO_CSV", "estado_validacion", "Pendiente / OK / Error", "NN"],
            ["FICHERO_CSV", "ruta_archivo", "Texto(260)", "NN"],
            ["REGLA_VALIDACIÓN", "id_regla", "Entero", "PK"],
            ["REGLA_VALIDACIÓN", "tipo_fichero", "Igual a FICHERO_CSV", "NN"],
            ["REGLA_VALIDACIÓN", "nombre_campo", "Texto(40)", "NN"],
            ["REGLA_VALIDACIÓN", "tipo_dato", "Texto / fecha / decimal", "NN"],
            ["REGLA_VALIDACIÓN", "longitud", "Entero", "Opcional"],
            ["REGLA_VALIDACIÓN", "obligatorio", "Booleano", "NN"],
            ["REGLA_VALIDACIÓN", "expresion_regla", "Texto(300)", "NN"],
            ["REGLA_VALIDACIÓN", "mensaje_error", "Texto(200)", "NN"],
            ["RESULTADO_VALIDACIÓN", "id_resultado", "Entero", "PK"],
            ["RESULTADO_VALIDACIÓN", "id_fichero", "Entero", "FK → FICHERO_CSV, NN"],
            ["RESULTADO_VALIDACIÓN", "id_regla", "Entero", "FK → REGLA_VALIDACIÓN, NN"],
            ["RESULTADO_VALIDACIÓN", "fila_afectada", "Entero", "Opcional"],
            ["RESULTADO_VALIDACIÓN", "resultado", "OK / Error / Advertencia", "NN"],
            ["RESULTADO_VALIDACIÓN", "detalle", "Texto(400)", "Opcional"],
            ["CARGUE", "id_cargue", "Entero", "PK"],
            ["CARGUE", "id_fichero", "Entero", "FK → FICHERO_CSV, NN"],
            ["CARGUE", "id_usuario_carga", "Entero", "FK → USUARIO, NN"],
            ["CARGUE", "id_usuario_aprueba", "Entero", "FK → USUARIO, NN"],
            ["CARGUE", "fecha_cargue", "Fecha-hora", "NN"],
            ["CARGUE", "medio_cargue", "FTP / IMPORTACION_12.5", "NN"],
            ["CARGUE", "resultado", "Exitoso / Error", "NN"],
            ["CARGUE", "log_sistema", "Texto largo", "Opcional"],
            ["CARGUE", "registro_sox", "Texto(40)", "Control AVH-263-F"],
        ],
        col_widths=[4.2, 4.2, 4.6, 3.5],
    )
    note_table(
        doc,
        "Los atributos id_usuario_carga e id_usuario_aprueba cubren el objetivo específico 3 y la separación de funciones. El hash del archivo origen permite detectar alteraciones entre la solicitud aprobada y el fichero generado.",
    )

    heading2(doc, "Atributos de reforma tarifaria")
    caption_table(doc, 12, "Atributos de DESCUENTO_TARIFA, MÍNIMA_DESPACHO y CARGO_MANEJO")
    add_table(
        doc,
        ["Entidad", "Atributo", "Tipo / dominio", "Restricción"],
        [
            ["DESCUENTO_TARIFA", "id_descuento", "Entero", "PK"],
            ["DESCUENTO_TARIFA", "id_cliente", "Entero", "FK, NN (agrupador)"],
            ["DESCUENTO_TARIFA", "id_centro", "Entero", "FK, NN"],
            ["DESCUENTO_TARIFA", "id_producto", "Entero", "FK, NN"],
            ["DESCUENTO_TARIFA", "id_punto_venta", "Entero", "FK, opcional"],
            ["DESCUENTO_TARIFA", "id_zona", "Entero", "FK, opcional"],
            ["DESCUENTO_TARIFA", "pais_origen / destino", "Texto(3)", "FK, opcional"],
            ["DESCUENTO_TARIFA", "cp_origen / destino", "Texto(6)", "Opcional"],
            ["DESCUENTO_TARIFA", "concepto_facturable", "Texto(40)", "NN"],
            ["DESCUENTO_TARIFA", "desde_kg / hasta_kg", "Decimal, coma, 3 decimales", "Opcional"],
            ["DESCUENTO_TARIFA", "descuento", "Decimal, 2 decimales", "NN"],
            ["DESCUENTO_TARIFA", "fecha_desde / hasta / baja", "Fecha dd/mm/aaaa", "desde NN"],
            ["MÍNIMA_DESPACHO", "id_minima", "Entero", "PK"],
            ["MÍNIMA_DESPACHO", "id_cliente / id_centro / id_producto", "Entero", "FK, NN"],
            ["MÍNIMA_DESPACHO", "rango_kg", "Decimal, 3 decimales", "NN"],
            ["MÍNIMA_DESPACHO", "minimo", "Decimal, 2 decimales", "NN"],
            ["MÍNIMA_DESPACHO", "baremo", "Entero 0–3 dígitos", "Opcional"],
            ["MÍNIMA_DESPACHO", "fecha_desde / hasta / baja", "Fecha", "desde NN"],
            ["CARGO_MANEJO", "id_cargo", "Entero", "PK"],
            ["CARGO_MANEJO", "id_cliente / id_centro / id_producto", "Entero", "FK, NN"],
            ["CARGO_MANEJO", "tasa_manejo", "Decimal, 2 decimales", "NN"],
            ["CARGO_MANEJO", "minimo", "Decimal, 2 decimales", "NN"],
            ["CARGO_MANEJO", "fecha_desde / hasta / baja", "Fecha", "desde NN"],
        ],
        col_widths=[4.2, 5.2, 4.2, 3.0],
    )
    note_table(
        doc,
        "La lógica de reforma exige tres registros cuando hay tarifa activa: baja de la actual, ajuste de fecha hasta al día anterior a la nueva vigencia, y registro de la tarifa nueva. Esas tres filas se persisten como ocurrencias de la misma entidad, diferenciadas por fecha_baja y fecha_hasta (Deprisa, 2026).",
    )

    # ===================== RELACIONES =====================
    heading1(doc, "Relaciones del Proyecto")
    add_body(
        doc,
        "Las relaciones expresan las reglas de negocio que el proceso manual hoy aplica en Excel y "
        "que el sistema automatizado debe garantizar por integridad referencial. La Tabla 13 resume "
        "nombre, entidades participantes, cardinalidad y semántica.",
    )
    caption_table(doc, 13, "Relaciones, cardinalidades y reglas de negocio")
    add_table(
        doc,
        ["Relación", "Entidades", "Cardinalidad", "Semántica"],
        [
            ["posee", "ROL – USUARIO", "1:N", "Un rol es asignado a muchos usuarios."],
            ["tiene", "CLIENTE – CENTRO", "1:N", "El centro pertenece a un agrupador existente."],
            ["posee", "CLIENTE – TARIFA", "1:N", "La tarifa especial o de reforma pertenece a un cliente; la general B2C puede no tener cliente dueño."],
            ["tiene", "TARIFA – VIGENCIA", "1:N", "Una tarifa tiene varias vigencias históricas; no se eliminan productos, se crean vigencias nuevas."],
            ["define", "VIGENCIA – BAREMO", "1:N", "Los baremos viven dentro de una vigencia y un producto."],
            ["desglosa", "BAREMO – DETALLE_BAREMO", "1:N", "Cada baremo se descompone en rangos hasta/importe/sobrepeso."],
            ["cubre", "VIGENCIA – RUTA", "1:N", "La malla de cobertura pertenece a una vigencia."],
            ["aplica en", "BAREMO – RUTA", "1:N", "Toda ruta apunta a un baremo; el FTP exige cargar ambos ficheros a la vez."],
            ["usa", "PRODUCTO – RUTA / BAREMO", "1:N", "Producto de 4 dígitos existente en el catálogo."],
            ["contiene", "PAÍS – CÓDIGO_POSTAL", "1:N", "El código postal pertenece a un país existente."],
            ["ubica", "PAÍS / CÓDIGO_POSTAL – RUTA", "1:N", "Origen y destino de la ruta."],
            ["origina", "CLIENTE – SOLICITUD", "1:N", "La solicitud B2B se asocia al cliente de la hoja tarifaria."],
            ["registra", "USUARIO – SOLICITUD", "1:N", "Quién recibe y deja trazabilidad del caso."],
            ["adjunta", "SOLICITUD – ARCHIVO_INSUMO", "1:N", "Una solicitud trae varios Excel."],
            ["genera", "SOLICITUD – FICHERO_CSV", "1:N", "De una solicitud salen, como mínimo, rutas y baremos."],
            ["se valida", "FICHERO_CSV – RESULTADO_VALIDACIÓN", "1:N", "Cada fichero produce hallazgos de reglas."],
            ["aplica", "REGLA_VALIDACIÓN – RESULTADO_VALIDACIÓN", "1:N", "El catálogo de lineamientos se reutiliza."],
            ["se carga", "FICHERO_CSV – CARGUE", "1:N", "Un fichero puede reintentarse si el log reporta error."],
            ["ejecuta / aprueba", "USUARIO – CARGUE", "1:N", "Separación de quien carga y quien aprueba."],
            ["aplica", "CLIENTE / CENTRO / PRODUCTO – reforma", "1:N", "Descuento, mínima y cargo se tasán sobre maestros existentes."],
        ],
        col_widths=[3.0, 4.2, 2.3, 7.0],
    )
    note_table(doc, "Elaboración propia. La cardinalidad 1:N es la predominante; no se identificaron relaciones N:M que requieran tabla asociativa adicional.")

    heading2(doc, "Reglas de integridad relevantes")
    add_body(
        doc,
        "Además de las llaves foráneas, el modelo incorpora restricciones que el proceso manual hoy "
        "verifica a ojo y que el sistema propuesto debe automatizar:",
    )
    add_body(
        doc,
        "1. Unicidad de ruta: no pueden existir dos rutas con la misma combinación de vigencia, "
        "producto, país origen, código postal origen, país destino y código postal destino.",
    )
    add_body(
        doc,
        "2. Dependencia ruta-baremo: no se autoriza un CARGUE de tipo RUTAS sin un FICHERO_CSV "
        "hermano de tipo BAREMOS en la misma SOLICITUD, ni a la inversa.",
    )
    add_body(
        doc,
        "3. Vigencia y producto: si ya existen productos en una vigencia, el cargue debe rechazarse "
        "o exigir una vigencia nueva, en lugar de borrar histórico.",
    )
    add_body(
        doc,
        "4. Fechas de reforma: fecha_hasta no puede ser inferior a fecha_desde; fecha_baja marca "
        "inactivación y conserva el histórico.",
    )
    add_body(
        doc,
        "5. Formato CSV: separador de listas igual a tubería (|), celdas en texto, sin fórmulas, "
        "sin encabezados en el archivo final, y nombres con patrón de fecha más consecutivo.",
    )
    add_body(
        doc,
        "6. Segregación de funciones: id_usuario_carga e id_usuario_aprueba no deberían coincidir "
        "en un mismo CARGUE, salvo rol administrador documentado.",
    )

    # ===================== DIAGRAMAS =====================
    heading1(doc, "Diagrama Entidad-Relación")
    add_body(
        doc,
        "Se presentan dos vistas complementarias. La Figura 1 usa la notación de Chen para enfatizar "
        "entidades (rectángulos), relaciones (rombos) y cardinalidades 1:N, que es la vista conceptual "
        "solicitada en Bases de Datos I. La Figura 2 usa notación Crow’s Foot para mostrar el modelo "
        "lógico con llaves primarias, foráneas y atributos principales, alineado con la "
        "implementación posterior del objetivo específico 1 (Elmasri y Navathe, 2016; Silberschatz et al., 2020).",
    )

    p = doc.add_paragraph()
    set_paragraph_format(p, align="center", first_line=False)
    run = p.add_run()
    run.add_picture(str(FIG_CHEN), width=Inches(6.5))
    caption_figure(
        doc,
        1,
        "Diagrama entidad-relación conceptual del sistema de automatización tarifaria, en notación de Chen. "
        "Las cajas verdes son maestros; las azules, el flujo de conversión y la reforma; los rombos dorados, las relaciones.",
    )

    p = doc.add_paragraph()
    set_paragraph_format(p, align="center", first_line=False)
    run = p.add_run()
    run.add_picture(str(FIG_CROW), width=Inches(6.5))
    caption_figure(
        doc,
        2,
        "Modelo lógico entidad-relación en notación Crow’s Foot, con claves y atributos principales. "
        "La flecha apunta al lado N de cada relación 1:N. Verde: maestros tarifarios; dorado: automatización; azul: reforma.",
    )

    heading2(doc, "Lectura del diagrama")
    add_body(
        doc,
        "El eje vertical del modelo es CLIENTE → TARIFA → VIGENCIA → (RUTA, BAREMO). Ese eje "
        "reproduce el maestro 1.2.1. El eje horizontal de automatización es SOLICITUD → "
        "ARCHIVO_INSUMO → FICHERO_CSV → (RESULTADO_VALIDACIÓN, CARGUE). Ese eje es la solución "
        "propuesta: los Excel dejan de ser el repositorio de verdad y pasan a ser insumos "
        "trazables. El módulo de IA no aparece como entidad porque no persiste identidad propia: "
        "es un componente de software que lee ARCHIVO_INSUMO, consulta REGLA_VALIDACIÓN y escribe "
        "FICHERO_CSV y RESULTADO_VALIDACIÓN (Sommerville, 2016).",
    )
    add_body(
        doc,
        "La reforma tarifaria se modeló en tres entidades especializadas y no en una sola tabla "
        "genérica, porque los ficheros de descuento, mínima y cargo tienen columnas distintas y se "
        "importan por tipos de información diferentes en el maestro 12.5. Forzar una única entidad "
        "habría generado nulos sistemáticos y habría debilitado las validaciones de longitud y tipo.",
    )

    # ===================== CONCLUSIONES =====================
    heading1(doc, "Conclusiones")
    add_body(
        doc,
        "El grupo de trabajo se mantuvo y se documentaron todos los ítems del formato de la entrega 1. "
        "Sobre ese proyecto se identificaron 23 entidades, sus atributos y 20 relaciones 1:N, "
        "materializadas en dos diagramas entidad-relación. El modelo cubre tanto el dominio tarifario "
        "(cliente, producto, tarifa, vigencia, ruta, baremo y reforma) como el flujo que se desea "
        "automatizar (solicitud, archivos Excel, CSV, reglas, validación, aprobación y cargue).",
    )
    add_body(
        doc,
        "Con este diseño queda satisfecho, en su componente de datos, el objetivo específico 1, y se "
        "deja la base estructural para los objetivos 2 y 3: el conversor asistido por IA tiene un "
        "catálogo de reglas persistido, y el programa de cargue tiene dónde registrar responsable, "
        "aprobador, archivo origen, fichero generado y resultado. La siguiente fase del proyecto "
        "consiste en traducir este modelo a un esquema SQL, implementar las restricciones de "
        "unicidad de ruta y de dependencia rutas-baremos, y construir los módulos de conversión y cargue.",
    )

    # ===================== REFERENCIAS =====================
    heading1(doc, "Referencias")

    refs = [
        "American Psychological Association. (2020). Publication manual of the American Psychological Association (7th ed.). https://doi.org/10.1037/0000165-000",
        "Chen, P. P.-S. (1976). The entity-relationship model—Toward a unified view of data. ACM Transactions on Database Systems, 1(1), 9–36. https://doi.org/10.1145/320434.320440",
        "Connolly, T. y Begg, C. (2015). Database systems: A practical approach to design, implementation, and management (6.ª ed.). Pearson.",
        "Deprisa. (2026). Parametrización tarifaria. Sistema Alertran Padua (Rev. 00) [Instructivo interno].",
        "Elmasri, R. y Navathe, S. B. (2016). Fundamentals of database systems (7.ª ed.). Pearson.",
        "Pressman, R. S. y Maxim, B. R. (2020). Software engineering: A practitioner’s approach (9.ª ed.). McGraw-Hill.",
        "Silberschatz, A., Korth, H. F. y Sudarshan, S. (2020). Database system concepts (7.ª ed.). McGraw-Hill.",
        "Sommerville, I. (2016). Software engineering (10.ª ed.). Pearson.",
        "Universidad Militar Nueva Granada. (s. f.). Formato de proyecto de actividad integradora [Documento institucional].",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 2.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-1.27)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(ref)
        set_run_font(run)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
